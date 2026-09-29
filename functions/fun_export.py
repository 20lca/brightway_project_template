"""Excel report export styled after Report example, AAK.dotx.

Reference: TableGrid headers #6DAACB, white text; body TT Commons Pro,
9 pt, #353B37; section rows #BFD2D0; thin table rules.
The style is embedded so exporting does not require the Word template.
"""
from math import ceil, floor, isfinite, log10
from numbers import Real

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER = "6DAACB"
SECTION = "BFD2D0"
TEXT = "353B37"
FONT = "TT Commons Pro"


def _value(value):
    if isinstance(value, (tuple, list, dict)):
        return str(value)
    if pd.isna(value):
        return None
    return value.item() if hasattr(value, "item") else value


def _number_format(value):
    """Show at least six significant digits without scientific notation."""
    if isinstance(value, int):
        return "#,##0"
    decimals = max(2, 5 - floor(log10(abs(value)))) if value and isfinite(value) else 2
    return "#,##0." + "00" + "#" * min(248, decimals - 2)


def _display_text(cell):
    """Estimate dimensions using the displayed decimal value, including zero."""
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, Real) and not isinstance(value, bool):
        decimals = len(cell.number_format.partition(".")[2])
        text = format(value, f",.{decimals}f")
        if decimals > 2:
            text = text.rstrip("0")
            if len(text.partition(".")[2]) < 2:
                text += "0" * (2 - len(text.partition(".")[2]))
        return text
    return str(value)


def _row(ws, values, *, header=False, section=False):
    ws.append([_value(value) for value in values])
    row = ws.max_row
    for cell in ws[row]:
        # Labels from databases are text, even if they begin with "=".
        if isinstance(cell.value, str):
            cell.data_type = "s"
        cell.font = Font(name=FONT, size=9, bold=header or section,
                         color="FFFFFF" if header else TEXT)
        cell.fill = PatternFill("solid", fgColor=HEADER if header else SECTION if section else "FFFFFF")
        cell.alignment = Alignment(vertical="top", wrap_text=True,
                                   horizontal="right" if isinstance(cell.value, Real) and not isinstance(cell.value, bool) else "left")
        cell.border = Border(bottom=Side(style="thin", color=TEXT))
        if isinstance(cell.value, Real) and not isinstance(cell.value, bool):
            cell.number_format = _number_format(cell.value)
    return row


def _finish(ws, *, filter_rows=True):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    merged_starts = {(r.min_row, r.min_col): r for r in ws.merged_cells.ranges}
    for column in ws.columns:
        longest = max(
            (max(map(len, _display_text(cell).split("\n")))
             for cell in column if (cell.row, cell.column) not in merged_starts),
            default=0,
        )
        dimension = ws.column_dimensions[get_column_letter(column[0].column)]
        dimension.width = min(48, max(6, longest + 2))
        dimension.bestFit = True
    for row in ws:
        lines = 1
        for cell in row:
            merged = merged_starts.get((cell.row, cell.column))
            columns = range(merged.min_col, merged.max_col + 1) if merged else [cell.column]
            width = sum(ws.column_dimensions[get_column_letter(c)].width for c in columns)
            lines = max(lines, sum(
                max(1, ceil(len(line) / max(1, width - 2)))
                for line in _display_text(cell).split("\n")
            ))
        ws.row_dimensions[row[0].row].height = lines * 12 + 3 if any(
            cell.value is not None for cell in row) else 6
    if filter_rows:
        ws.auto_filter.ref = ws.dimensions
    ws.print_title_rows = "1:1"
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3 if ws.max_column > 10 else ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = ws.dimensions


def _table(wb, name, frame):
    ws = wb.create_sheet(name)
    _row(ws, list(frame.columns), header=True)
    for values in frame.itertuples(index=False, name=None):
        _row(ws, values)
    _finish(ws)
    return ws



def _inventory_sheet(wb, activities):
    """Write unfiltered direct exchanges on their stored production basis."""
    ws = wb.create_sheet("Life Cycle Inventory")
    _row(ws, ["Exchange", "Unit", "Amount", "Comments"], header=True)

    def band(text):
        row = _row(ws, [text, None, None, None], section=True)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)

    for index, (label, activity) in enumerate(activities.items()):
        if index:
            ws.append([None] * 4)
            _row(ws, ["Exchange", "Unit", "Amount", "Comments"], header=True)
        for category, exchanges in (
                ("Production / reference basis", activity.production()),
                ("Technosphere", activity.technosphere()),
                ("Biosphere", activity.biosphere())):
            exchanges = list(exchanges)
            if category == 'Biosphere' and not exchanges:
                continue
            if category != 'Production / reference basis':
                band(category)
            count = 0
            for exchange in exchanges:
                node = exchange.input
                name = node.get("name", "(Unnamed exchange)")
                if category == "Production / reference basis":
                    name = node.get("reference product") or name
                if category == "Biosphere":
                    compartments = node.get("categories") or ()
                    if isinstance(compartments, str):
                        compartments = [compartments]
                    detail = " / ".join(str(x) for x in compartments)
                else:
                    detail = node.get("location", "")
                if detail:
                    name += f" ({detail})"
                _row(ws, [name, exchange.get("unit") or node.get("unit"),
                          exchange["amount"], exchange.get("comment", "")])
                count += 1
            if not count:
                message = ("No explicit production exchange; no basis inferred"
                           if category == "Production / reference basis"
                           else "No exchanges")
                _row(ws, [message, None, None, None])
    if not activities:
        _row(ws, ["No activities supplied", None, None, None])
    _finish(ws, filter_rows=False)
    return ws


def export_workflow_excel(path, *, activity_selection, scores, scores_wide,
                          contributions, score_checks, activities=None):
    """Export existing analysis tables; preserve values and all retained branches.

    Contribution overview has green section rows and marked total rows.
    Direct flows and cutoff-only other rows are supplied by the workflow.
    Contributions data is grouped by studied process, retaining every raw column.
    Optional activities supply direct inventories on their stored production basis.
    This does not calculate an aggregated upstream life-cycle inventory.
    """
    if not contributions.empty and "contribution_type" not in contributions:
        raise ValueError("Recalculate contributions to separate direct impacts from cutoff omissions.")
    wb = Workbook()
    wb.remove(wb.active)
    _table(wb, "Activities", activity_selection)
    if activities is not None:
        _inventory_sheet(wb, activities)
    _table(wb, "LCA overview", scores_wide)
    _table(wb, "LCA scores", scores)
    ws = wb.create_sheet("Contribution overview")
    columns = [c for c in ("name", "location", "amount", "unit", "score", "fraction") if c in contributions]
    labels = {"name": "Contributor", "location": "Location", "amount": "Amount",
              "unit": "Unit", "score": "Impact", "fraction": "Share of total"}
    _row(ws, ["Row type"] + [labels[c] for c in columns], header=True)
    for (activity, category), group in contributions.groupby(
            ["activity_label", "category_key"], sort=False, dropna=False):
        unit = group["method_unit"].iloc[0]
        reference = group["reference_unit"].iloc[0]
        row = _row(ws, [f"{activity} | {category} | Impact: {unit} | Reference: 1 {reference}"]
                   + [None] * len(columns), section=True)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(columns) + 1)
        for _, item in group.iterrows():
            root = pd.isna(item["parent"])
            row = _row(ws, ["Total (do not sum)" if root else item["contribution_type"]] + [item[c] for c in columns])
            if root:
                for cell in ws[row]:
                    cell.font = Font(name=FONT, size=9, bold=True, color=TEXT)
    _finish(ws, filter_rows=False)
    ws = wb.create_sheet("Contributions data")
    if contributions.empty:
        _row(ws, list(contributions.columns), header=True)
    else:
        for index, (activity, group) in enumerate(contributions.groupby(
                "activity_label", sort=False, dropna=False)):
            if index:
                ws.append([None] * len(contributions.columns))
            row = _row(ws, [f"Studied process: {activity}"]
                       + [None] * (len(contributions.columns) - 1), section=True)
            ws.merge_cells(start_row=row, start_column=1, end_row=row,
                           end_column=len(contributions.columns))
            _row(ws, list(contributions.columns), header=True)
            for values in group.itertuples(index=False, name=None):
                _row(ws, values)
    _finish(ws, filter_rows=False)
    ws.freeze_panes = "A3" if not contributions.empty else "A2"
    ws.print_title_rows = "1:2" if not contributions.empty else "1:1"
    _table(wb, "Score checks", score_checks)
    notes = pd.DataFrame({"Notes": [
        "Table style adapted from Report example, AAK.dotx (TableGrid and green section rows).",
        "LCIA results and contributions are for one reference-product unit. Impact units are in category labels or method_unit.",
        "Contribution totals and input rows must not be added together.",
        "Inputs include upstream impacts. Direct rows show characterized biosphere flows of the studied process. Other contains only inputs omitted by CUTOFF. Shares are undefined when the total is zero.",
        "Contribution tables retain all branches passing CUTOFF; MAX_CONTRIBUTORS limits figures only.",
        "Input and direct-flow amounts are scaled to one reference-product unit; their units are shown in the unit column.",
        "Life Cycle Inventory, when included, lists direct process exchanges grouped by type. Production exchanges show the stored basis. Amounts and signs are unchanged; zero and negative exchanges are retained without cutoff, normalization or upstream aggregation. Comments come from exchange metadata.",
    ]})
    _table(wb, "Read me", notes)
    wb.save(path)
    return path
